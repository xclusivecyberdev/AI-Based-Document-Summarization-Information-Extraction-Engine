"""
Document Processing Module
Handles PDF, DOCX, images, and scanned documents
"""
import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import io

import PyPDF2
import pdfplumber
from docx import Document
from PIL import Image
import chardet
from loguru import logger


class DocumentProcessor:
    """Process various document formats and extract text"""

    SUPPORTED_FORMATS = {
        'pdf': ['.pdf'],
        'docx': ['.docx', '.doc'],
        'image': ['.png', '.jpg', '.jpeg', '.tiff', '.bmp', '.gif'],
        'text': ['.txt', '.md', '.csv']
    }

    def __init__(self):
        self.supported_extensions = []
        for formats in self.SUPPORTED_FORMATS.values():
            self.supported_extensions.extend(formats)

    def process_document(self, file_path: str) -> Dict:
        """
        Process a document and extract text and metadata

        Args:
            file_path: Path to the document

        Returns:
            Dictionary containing extracted text, metadata, and structure
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        extension = file_path.suffix.lower()

        if extension not in self.supported_extensions:
            raise ValueError(f"Unsupported file format: {extension}")

        logger.info(f"Processing document: {file_path.name}")

        # Route to appropriate processor
        if extension in self.SUPPORTED_FORMATS['pdf']:
            return self._process_pdf(file_path)
        elif extension in self.SUPPORTED_FORMATS['docx']:
            return self._process_docx(file_path)
        elif extension in self.SUPPORTED_FORMATS['image']:
            return self._process_image(file_path)
        elif extension in self.SUPPORTED_FORMATS['text']:
            return self._process_text(file_path)
        else:
            raise ValueError(f"No processor available for: {extension}")

    def _process_pdf(self, file_path: Path) -> Dict:
        """Extract text and metadata from PDF files"""
        result = {
            'text': '',
            'pages': [],
            'metadata': {},
            'tables': [],
            'images': [],
            'structure': {}
        }

        try:
            # Extract text using PyPDF2
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)

                # Metadata
                if pdf_reader.metadata:
                    result['metadata'] = {
                        'title': pdf_reader.metadata.get('/Title', ''),
                        'author': pdf_reader.metadata.get('/Author', ''),
                        'subject': pdf_reader.metadata.get('/Subject', ''),
                        'creator': pdf_reader.metadata.get('/Creator', ''),
                        'producer': pdf_reader.metadata.get('/Producer', ''),
                    }

                result['metadata']['num_pages'] = len(pdf_reader.pages)

                # Extract text from each page
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    page_text = page.extract_text()
                    result['pages'].append({
                        'page_number': page_num,
                        'text': page_text
                    })
                    result['text'] += page_text + '\n\n'

            # Use pdfplumber for better table extraction
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract tables
                    tables = page.extract_tables()
                    if tables:
                        for table_idx, table in enumerate(tables):
                            result['tables'].append({
                                'page': page_num,
                                'table_index': table_idx,
                                'data': table
                            })

        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}")
            raise

        result['file_type'] = 'pdf'
        result['file_name'] = file_path.name
        return result

    def _process_docx(self, file_path: Path) -> Dict:
        """Extract text and metadata from DOCX files"""
        result = {
            'text': '',
            'paragraphs': [],
            'tables': [],
            'metadata': {},
            'structure': {}
        }

        try:
            doc = Document(file_path)

            # Metadata
            core_props = doc.core_properties
            result['metadata'] = {
                'title': core_props.title or '',
                'author': core_props.author or '',
                'subject': core_props.subject or '',
                'created': str(core_props.created) if core_props.created else '',
                'modified': str(core_props.modified) if core_props.modified else '',
            }

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    result['paragraphs'].append({
                        'text': para.text,
                        'style': para.style.name if para.style else 'Normal'
                    })
                    result['text'] += para.text + '\n'

            # Extract tables
            for table_idx, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)

                result['tables'].append({
                    'table_index': table_idx,
                    'data': table_data
                })

        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}")
            raise

        result['file_type'] = 'docx'
        result['file_name'] = file_path.name
        return result

    def _process_image(self, file_path: Path) -> Dict:
        """Process image files (will use OCR in next module)"""
        result = {
            'text': '',
            'metadata': {},
            'image_info': {},
            'requires_ocr': True
        }

        try:
            with Image.open(file_path) as img:
                result['image_info'] = {
                    'format': img.format,
                    'mode': img.mode,
                    'size': img.size,
                    'width': img.width,
                    'height': img.height
                }

                # Get EXIF data if available
                if hasattr(img, '_getexif') and img._getexif():
                    result['metadata']['exif'] = img._getexif()

        except Exception as e:
            logger.error(f"Error processing image: {str(e)}")
            raise

        result['file_type'] = 'image'
        result['file_name'] = file_path.name
        result['text'] = '[Image requires OCR processing]'

        return result

    def _process_text(self, file_path: Path) -> Dict:
        """Process plain text files"""
        result = {
            'text': '',
            'metadata': {},
            'encoding': 'utf-8'
        }

        try:
            # Detect encoding
            with open(file_path, 'rb') as file:
                raw_data = file.read()
                detected = chardet.detect(raw_data)
                encoding = detected['encoding'] or 'utf-8'
                result['encoding'] = encoding

            # Read text
            with open(file_path, 'r', encoding=encoding) as file:
                result['text'] = file.read()

            result['metadata'] = {
                'size_bytes': file_path.stat().st_size,
                'lines': len(result['text'].split('\n'))
            }

        except Exception as e:
            logger.error(f"Error processing text file: {str(e)}")
            raise

        result['file_type'] = 'text'
        result['file_name'] = file_path.name

        return result

    def is_supported(self, file_path: str) -> bool:
        """Check if file format is supported"""
        extension = Path(file_path).suffix.lower()
        return extension in self.supported_extensions

    def get_file_type(self, file_path: str) -> Optional[str]:
        """Get the type category of a file"""
        extension = Path(file_path).suffix.lower()

        for file_type, extensions in self.SUPPORTED_FORMATS.items():
            if extension in extensions:
                return file_type

        return None
