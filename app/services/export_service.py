"""
Export Service for DOCX and PDF
"""
from pathlib import Path
from typing import Dict, Optional
from datetime import datetime
import io

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from loguru import logger


class ExportService:
    """Export summaries and analysis to DOCX and PDF"""

    def export_to_docx(
        self,
        content: Dict,
        output_path: str,
        include_metadata: bool = True
    ) -> str:
        """
        Export content to DOCX format

        Args:
            content: Dictionary with summary and metadata
            output_path: Path to save the document
            include_metadata: Include metadata in document

        Returns:
            Path to saved document
        """
        logger.info(f"Exporting to DOCX: {output_path}")

        doc = Document()

        # Title
        title = doc.add_heading('Document Summary Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        if include_metadata:
            doc.add_heading('Metadata', level=1)

            metadata_table = doc.add_table(rows=0, cols=2)
            metadata_table.style = 'Light Grid Accent 1'

            metadata = content.get('metadata', {})

            if 'created_at' in metadata or 'timestamp' in metadata:
                row = metadata_table.add_row()
                row.cells[0].text = 'Generated'
                row.cells[1].text = str(metadata.get('created_at') or metadata.get('timestamp', datetime.now()))

            if 'method' in content:
                row = metadata_table.add_row()
                row.cells[0].text = 'Method'
                row.cells[1].text = content['method']

            if 'original_length' in content:
                row = metadata_table.add_row()
                row.cells[0].text = 'Original Length'
                row.cells[1].text = f"{content['original_length']} characters"

            if 'compression_ratio' in content:
                row = metadata_table.add_row()
                row.cells[0].text = 'Compression Ratio'
                row.cells[1].text = f"{content['compression_ratio']:.2%}"

            doc.add_paragraph()

        # Summary
        doc.add_heading('Summary', level=1)

        summary_text = content.get('summary', content.get('summary_text', ''))
        summary_para = doc.add_paragraph(summary_text)

        # Bullets if available
        if 'bullets' in content and content['bullets']:
            doc.add_heading('Key Points', level=1)

            for bullet in content['bullets']:
                doc.add_paragraph(bullet.replace('• ', ''), style='List Bullet')

        # Keywords if available
        if 'keywords' in content:
            doc.add_heading('Keywords', level=1)

            keywords = content['keywords']
            if isinstance(keywords, dict):
                for method, kw_list in keywords.items():
                    doc.add_heading(method.upper(), level=2)

                    if isinstance(kw_list, list):
                        for kw in kw_list[:10]:
                            if isinstance(kw, dict):
                                kw_text = kw.get('keyword', str(kw))
                            else:
                                kw_text = str(kw)

                            doc.add_paragraph(kw_text, style='List Bullet')

        # Entities if available
        if 'entities' in content:
            doc.add_heading('Named Entities', level=1)

            entities = content['entities']
            if isinstance(entities, dict):
                for entity_type, entity_list in entities.items():
                    if entity_list:
                        doc.add_heading(entity_type, level=2)

                        for entity in entity_list[:20]:
                            if isinstance(entity, dict):
                                entity_text = entity.get('text', entity.get('name', str(entity)))
                            else:
                                entity_text = str(entity)

                            doc.add_paragraph(entity_text, style='List Bullet')

        # Save document
        doc.save(output_path)

        logger.info(f"DOCX exported successfully: {output_path}")

        return output_path

    def export_to_pdf(
        self,
        content: Dict,
        output_path: str,
        include_metadata: bool = True
    ) -> str:
        """
        Export content to PDF format

        Args:
            content: Dictionary with summary and metadata
            output_path: Path to save the PDF
            include_metadata: Include metadata in PDF

        Returns:
            Path to saved PDF
        """
        logger.info(f"Exporting to PDF: {output_path}")

        # Create PDF
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=18
        )

        # Container for PDF elements
        story = []

        # Styles
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor='#1f4788',
            spaceAfter=30,
            alignment=1  # Center
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor='#1f4788',
            spaceAfter=12,
            spaceBefore=12
        )

        # Title
        story.append(Paragraph("Document Summary Report", title_style))
        story.append(Spacer(1, 0.2 * inch))

        # Metadata
        if include_metadata:
            story.append(Paragraph("Metadata", heading_style))

            metadata = content.get('metadata', {})

            if 'created_at' in metadata or 'timestamp' in metadata:
                timestamp = metadata.get('created_at') or metadata.get('timestamp', datetime.now())
                story.append(Paragraph(f"<b>Generated:</b> {timestamp}", styles['Normal']))

            if 'method' in content:
                story.append(Paragraph(f"<b>Method:</b> {content['method']}", styles['Normal']))

            if 'original_length' in content:
                story.append(Paragraph(
                    f"<b>Original Length:</b> {content['original_length']} characters",
                    styles['Normal']
                ))

            if 'compression_ratio' in content:
                ratio = content['compression_ratio']
                story.append(Paragraph(
                    f"<b>Compression Ratio:</b> {ratio:.2%}",
                    styles['Normal']
                ))

            story.append(Spacer(1, 0.2 * inch))

        # Summary
        story.append(Paragraph("Summary", heading_style))

        summary_text = content.get('summary', content.get('summary_text', ''))

        # Split into paragraphs
        paragraphs = summary_text.split('\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para, styles['Normal']))
                story.append(Spacer(1, 0.1 * inch))

        # Bullets if available
        if 'bullets' in content and content['bullets']:
            story.append(Spacer(1, 0.2 * inch))
            story.append(Paragraph("Key Points", heading_style))

            for bullet in content['bullets']:
                bullet_text = bullet.replace('• ', '')
                story.append(Paragraph(f"• {bullet_text}", styles['Normal']))

        # Build PDF
        doc.build(story)

        logger.info(f"PDF exported successfully: {output_path}")

        return output_path

    def export_analysis_to_docx(
        self,
        analysis: Dict,
        output_path: str
    ) -> str:
        """Export NLP analysis to DOCX"""
        logger.info(f"Exporting analysis to DOCX: {output_path}")

        doc = Document()

        # Title
        title = doc.add_heading('NLP Analysis Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Statistics
        doc.add_heading('Document Statistics', level=1)

        if 'statistics' in analysis:
            stats = analysis['statistics']
            stats_table = doc.add_table(rows=0, cols=2)
            stats_table.style = 'Light Grid Accent 1'

            for key, value in stats.items():
                row = stats_table.add_row()
                row.cells[0].text = key.replace('_', ' ').title()

                if isinstance(value, float):
                    row.cells[1].text = f"{value:.2f}"
                else:
                    row.cells[1].text = str(value)

        # Entities
        if 'entities' in analysis and analysis['entities']:
            doc.add_paragraph()
            doc.add_heading('Named Entities', level=1)

            for entity_type, entities in analysis['entities'].items():
                if entities:
                    doc.add_heading(entity_type, level=2)

                    for entity in entities[:20]:
                        text = entity.get('text', entity.get('name', str(entity)))
                        doc.add_paragraph(text, style='List Bullet')

        # Keywords
        if 'keywords' in analysis and analysis['keywords']:
            doc.add_paragraph()
            doc.add_heading('Keywords', level=1)

            for method, keywords in analysis['keywords'].items():
                if keywords:
                    doc.add_heading(method.upper(), level=2)

                    for kw in keywords[:15]:
                        if isinstance(kw, dict):
                            kw_text = kw.get('keyword', str(kw))
                        else:
                            kw_text = str(kw)

                        doc.add_paragraph(kw_text, style='List Bullet')

        doc.save(output_path)

        logger.info(f"Analysis DOCX exported successfully: {output_path}")

        return output_path
